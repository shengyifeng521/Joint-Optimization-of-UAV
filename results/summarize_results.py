#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""批量汇总已有实验结果 JSON，匹配原始实例，导出 XLSX。

安装依赖：python3 -m pip install openpyxl
运行示例：python3 summarize_results.py data --instance-dir instances -o summary.xlsx
"""

import argparse
import json
import math
import re
import sys
from pathlib import Path


HEADERS = [
    "实例名", "视点数", "采样点数", "S1 obj", "S2 obj", "total obj",
    "runtime", "选择的视点数量", "拍照数量", "求解状态status",
]
RESULT_KEYS = {
    "quality_norm", "photo_norm", "energy_norm", "status",
    "selected_vps", "n_viewpoints", "photos",
}


def read_json(filename):
    with filename.open("r", encoding="utf-8-sig") as stream:
        return json.load(stream)


def is_number(value):
    return type(value) in (int, float) and math.isfinite(value)


def is_count(value):
    return is_number(value) and value >= 0 and int(value) == value


def instance_name(filename):
    return re.sub(r"^(joint_|two_step_|sga_)", "", filename.stem, flags=re.I)


def natural_key(filename):
    return [
        (1, int(part)) if part.isdigit() else (0, part.casefold())
        for part in re.split(r"(\d+)", str(filename))
    ]


def input_count(data, array_key, metadata_key, warn):
    metadata = data.get("metadata")
    declared = metadata.get(metadata_key) if isinstance(metadata, dict) else None
    values = data.get(array_key)
    if isinstance(values, list):
        if is_count(declared) and declared != len(values):
            warn(f"{metadata_key}={declared} 与数组长度 {len(values)} 不一致，采用数组长度")
        return len(values)
    return int(declared) if is_count(declared) else None


def locate_instance(name, result_file, instance_dir=None):
    # 显式指定时只查该目录，避免匹配其他规模的同名实例。
    if instance_dir is not None:
        filename = instance_dir / f"{name}.json"
        return filename, read_json(filename)
    for parent in result_file.parents:
        filename = parent / "data" / f"{name}.json"
        try:
            data = read_json(filename)
        except FileNotFoundError:
            continue
        if isinstance(data, dict) and (
            isinstance(data.get("metadata"), dict)
            or isinstance(data.get("view_points"), list)
        ):
            return filename, data
    return None, None


def collect_rows(input_dir, instance_dir=None, recursive=False):
    """返回数据行和警告；缺失值使用 None，真实的 0 保留。"""
    input_dir = Path(input_dir).resolve()
    instance_dir = Path(instance_dir).resolve() if instance_dir is not None else None
    if not input_dir.is_dir():
        raise ValueError(f"结果目录不存在：{input_dir}")
    if instance_dir is not None and not instance_dir.is_dir():
        raise ValueError(f"原始实例目录不存在：{instance_dir}")
    candidates = input_dir.rglob("*") if recursive else input_dir.iterdir()
    files = sorted(
        (p for p in candidates if p.is_file() and p.suffix.lower() == ".json"),
        key=natural_key,
    )
    rows, warnings, sources = [], [], []
    skipped = 0
    for filename in files:
        relative = filename.relative_to(input_dir)

        def warn(message):
            warnings.append(f"{relative}: {message}")

        try:
            result = read_json(filename)
        except (OSError, ValueError) as error:
            warn(f"读取失败，已跳过：{error}")
            skipped += 1
            continue
        if not isinstance(result, dict) or not RESULT_KEYS.intersection(result):
            warn("不是实验结果 JSON，已跳过")
            skipped += 1
            continue
        name = instance_name(filename)
        display_name = (relative.parent / name).as_posix()
        n_views = n_samples = None
        try:
            source, data = locate_instance(name, filename, instance_dir)
            if isinstance(data, dict):
                sources.append(str(source))
                n_views = input_count(data, "view_points", "num_view_points", warn)
                n_samples = input_count(data, "sample_points", "num_samples_points", warn)
            if n_views is None or n_samples is None:
                warn("缺少原始实例或数量字段，对应项留空；可用 --instance-dir 指定原始实例目录")
        except (OSError, ValueError) as error:
            warn(f"读取原始实例失败，视点数/采样点数留空：{error}")

        def metric(key):
            value = result.get(key)
            if is_number(value):
                return value
            warn(f"缺少或无效的 {key}，对应值留空")
            return None

        quality = metric("quality_norm")
        photo = metric("photo_norm")
        s1 = -quality + photo if quality is not None and photo is not None else None
        s2 = metric("energy_norm")
        total = s1 + s2 if s1 is not None and s2 is not None else None
        runtime = metric("runtime")  # 使用 JSON 中的总 runtime。

        chosen = result.get("n_viewpoints")
        chosen = int(chosen) if is_count(chosen) else None
        for key in ("selected_vps", "selected_vp_ids"):
            if isinstance(result.get(key), list):
                chosen = len(result[key])
                break
        declared = result.get("n_viewpoints")
        if chosen is not None and is_count(declared) and chosen != declared:
            warn("选择的视点列表与 n_viewpoints 不一致，采用列表长度")
        photos = result.get("photos")
        photos = int(photos) if is_count(photos) else None
        if photos is None:
            for key in ("selected_pairs", "selected_photos_readable"):
                if isinstance(result.get(key), list):
                    photos = len(result[key])
                    break
        pairs = result.get("selected_pairs")
        if isinstance(pairs, list) and photos != len(pairs):
            warn("photos 与 selected_pairs 长度不一致，采用 photos")
        if chosen is None:
            warn("缺少选择的视点数量，留空")
        if photos is None:
            warn("缺少拍照数量，留空")
        status = result.get("status")
        if not (isinstance(status, (str, bool)) or is_number(status)):
            status = None
            warn("缺少或无效的求解状态 status，留空")
        if total is not None and is_number(result.get("obj")) and not math.isclose(
            total, result["obj"], rel_tol=1e-8, abs_tol=1e-8
        ):
            warn(f"按公式计算的 total obj={total} 与原 obj={result['obj']} 不一致，采用指定公式")
        rows.append([
            display_name, n_views, n_samples, s1, s2, total,
            runtime, chosen, photos, status,
        ])
        sources.append(str(filename))
    return {
        "headers": HEADERS, "rows": rows, "warnings": warnings,
        "sources": list(dict.fromkeys(sources)), "scanned": len(files), "skipped": skipped,
    }


def write_xlsx(rows, output):
    if not rows:
        raise ValueError("没有有效实验结果，不生成空汇总表")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError("缺少 openpyxl，请先运行：python3 -m pip install openpyxl") from error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "实验结果汇总"
    sheet.sheet_view.showGridLines = False
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows():
        for cell in row:
            # 文本按文本存储，避免以 = 开头的实例名被当成公式。
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=11, color="202C3A")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row[0].row].height = 24
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="223F5D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    for column, width in zip("ABCDEFGHIJ", [24, 12, 12, 16, 16, 16, 16, 21, 14, 23]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (2, 3, 8, 9):
                cell.number_format = "#,##0"
            elif cell.column in (4, 5, 6):
                cell.number_format = "0.000000"
            elif cell.column == 7:
                cell.number_format = "0.000"
    table = Table(displayName="ExperimentResults", ref=f"A1:J{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_dir", nargs="?", default="data", help="结果 JSON 目录，默认 data")
    parser.add_argument("--instance-dir", type=Path, help="原始实例 JSON 目录，用于读取视点数和采样点数")
    parser.add_argument("-o", "--output", type=Path, default=Path("summary.xlsx"))
    parser.add_argument("-r", "--recursive", action="store_true", help="递归遍历结果子目录")
    parser.add_argument("--report", type=Path, help="可选：保存来源和异常记录为 JSON")
    args = parser.parse_args()
    if args.output.suffix.lower() != ".xlsx":
        parser.error("输出扩展名必须为 .xlsx")
    if args.report and args.report.resolve() == args.output.resolve():
        parser.error("报告和 XLSX 路径必须不同")
    try:
        report = collect_rows(args.input_dir, args.instance_dir, args.recursive)
        for warning in report["warnings"]:
            print(f"警告：{warning}", file=sys.stderr)
        write_xlsx(report["rows"], args.output)
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"扫描 {report['scanned']} 个 JSON，汇总 {len(report['rows'])} 条，"
              f"跳过 {report['skipped']} 个，警告 {len(report['warnings'])} 条。")
        print(f"已生成：{args.output.resolve()}")
    except (OSError, ValueError, RuntimeError) as error:
        print(f"错误：{error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
